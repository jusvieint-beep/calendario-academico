"""
Genera la plantilla oficial de Excel del Calendario Académico.

Uso:
    pip install openpyxl
    python scripts/generar_plantilla.py

Crea public/plantilla-calendario.xlsx, que es el archivo que el panel
administrativo ofrece en "Descargar plantilla vacía".

Si cambias columnas aquí, cambia también src/lib/excel/columns.ts.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = 1000  # filas preparadas con formato y validaciones
OUT = Path(__file__).resolve().parent.parent / "public" / "plantilla-calendario.xlsx"

INK = "211F18"
DIM = "6B6858"
HEAD_FILL = PatternFill("solid", fgColor="1B1B1B")
ACCENT = "EFF422"
SES_FILL = PatternFill("solid", fgColor="DCE6FB")
ADI_FILL = PatternFill("solid", fgColor="DDF3F7")
DUD_FILL = PatternFill("solid", fgColor="FBF0D2")
ENT_FILL = PatternFill("solid", fgColor="FDDCDF")
REC_FILL = PatternFill("solid", fgColor="FBE1EE")
QUIZ_FILL = PatternFill("solid", fgColor="DDF3E6")
FORO_FILL = PatternFill("solid", fgColor="EEE4FA")
SOFT_FILL = PatternFill("solid", fgColor="F2F0E6")
thin = Side(style="thin", color="DDD9C8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# (columna, ancho, ayuda en la celda)
COLUMNS = [
    ("ID_EVENTO", 13, "Déjalo vacío en filas nuevas: el sistema asigna el ID. No cambies los ID existentes."),
    ("TIPO", 20, "Elige de la lista: SESION_ZAJUNA, SESION_ADICIONAL, DUDAS, GRABACION, ENTREGA, CUESTIONARIO o FORO."),
    ("NOMBRE", 34, "Nombre que verán los estudiantes. Máximo 150 caracteres."),
    ("FECHA", 13, "Formato dd/mm/aaaa. Ejemplo: 05/10/2026."),
    ("HORA_INICIO", 13, "Formato 24 h (08:00, 14:30). Obligatoria en sesiones y dudas. En entregas es la hora límite."),
    ("HORA_FIN", 11, "Solo para sesiones y dudas. Debe ser posterior a la hora de inicio."),
    ("LINK", 42, "Enlace de la sesión, del chat de dudas o del recurso (https://...). Opcional."),
    ("DESCRIPCION", 46, "Información adicional. Opcional. Máximo 2.000 caracteres."),
    ("INSTRUCTOR", 28, "Persona que dicta la sesión. Opcional. Para SESION_ZAJUNA, SESION_ADICIONAL, DUDAS y GRABACION."),
]


def build_calendar_sheet(ws):
    ws.title = "Calendario"
    for idx, (name, width, _help) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=ACCENT))
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ROWS + 1}"

    for r in range(2, ROWS + 2):
        ws.cell(row=r, column=1).number_format = "@"
        ws.cell(row=r, column=4).number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=5).number_format = "hh:mm"
        ws.cell(row=r, column=6).number_format = "hh:mm"
        ws.cell(row=r, column=8).alignment = Alignment(wrap_text=False)

    last = ROWS + 1
    tipo = DataValidation(type="list", formula1='"SESION_ZAJUNA,SESION_ADICIONAL,DUDAS,GRABACION,ENTREGA,CUESTIONARIO,FORO"', allow_blank=True,
                          showErrorMessage=True, errorTitle="Tipo no válido",
                          error="Elige SESION_ZAJUNA, SESION_ADICIONAL, DUDAS, GRABACION, ENTREGA, CUESTIONARIO o FORO.",
                          showInputMessage=True, promptTitle="TIPO", prompt=COLUMNS[1][2])
    tipo.add(f"B2:B{last}")

    fecha = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)",
                           allow_blank=True, showErrorMessage=True, errorTitle="Fecha no válida",
                           error="Escribe la fecha como dd/mm/aaaa, por ejemplo 05/10/2026.",
                           showInputMessage=True, promptTitle="FECHA", prompt=COLUMNS[3][2])
    fecha.add(f"D2:D{last}")

    hora = DataValidation(type="time", operator="between", formula1="TIME(0,0,0)", formula2="TIME(23,59,59)",
                          allow_blank=True, showErrorMessage=True, errorTitle="Hora no válida",
                          error="Escribe la hora en formato 24 h, por ejemplo 08:00 o 14:30.",
                          showInputMessage=True, promptTitle="HORA", prompt=COLUMNS[4][2])
    hora.add(f"E2:F{last}")

    nombre = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="150", allow_blank=True,
                            showErrorMessage=True, errorTitle="Nombre muy largo",
                            error="Máximo 150 caracteres.", showInputMessage=True,
                            promptTitle="NOMBRE", prompt=COLUMNS[2][2])
    nombre.add(f"C2:C{last}")

    idv = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showInputMessage=True,
                         promptTitle="ID_EVENTO", prompt=COLUMNS[0][2])
    idv.add(f"A2:A{last}")

    link = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showInputMessage=True,
                          promptTitle="LINK", prompt=COLUMNS[6][2])
    link.add(f"G2:G{last}")
    desc = DataValidation(type="custom", formula1="TRUE", allow_blank=True, showInputMessage=True,
                          promptTitle="DESCRIPCION", prompt=COLUMNS[7][2])
    desc.add(f"H2:H{last}")
    inst = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="120", allow_blank=True,
                          showErrorMessage=True, errorTitle="Nombre muy largo", error="Máximo 120 caracteres.",
                          showInputMessage=True, promptTitle="INSTRUCTOR", prompt=COLUMNS[8][2])
    inst.add(f"I2:I{last}")

    for dv in (tipo, fecha, hora, nombre, idv, link, desc, inst):
        ws.add_data_validation(dv)


def build_instructions_sheet(wb):
    ws = wb.create_sheet("Instrucciones")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    widths = {"A": 3, "B": 16, "C": 14, "D": 26, "E": 18, "F": 60}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    row = 2

    def title(text, size=16):
        nonlocal row
        c = ws.cell(row=row, column=2, value=text)
        c.font = Font(bold=True, size=size, color=INK)
        row += 1

    def para(text, bold=False, color=DIM, height=None):
        nonlocal row
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        c = ws.cell(row=row, column=2, value=text)
        c.font = Font(size=11, bold=bold, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if height:
            ws.row_dimensions[row].height = height
        row += 1

    def table(headers, rows, fills=None):
        nonlocal row
        for i, h in enumerate(headers):
            c = ws.cell(row=row, column=2 + i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
        row += 1
        for r_i, data in enumerate(rows):
            for i, v in enumerate(data):
                c = ws.cell(row=row, column=2 + i, value=v)
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.font = Font(color=INK)
                if fills and fills[r_i]:
                    c.fill = fills[r_i]
            row += 1

    title("Calendario Académico · Cómo llenar esta plantilla", 18)
    para("Llena la hoja «Calendario» con TODAS las sesiones, espacios de dudas, grabaciones, entregas, cuestionarios y foros. "
         "Una fila por actividad. No cambies los nombres de las columnas ni el nombre de la hoja.", height=32)
    row += 1

    title("Lo más importante")
    para("1. El archivo representa el calendario COMPLETO. Lo que no esté en el archivo se eliminará del calendario "
         "al confirmar la importación (antes verás una vista previa y deberás confirmar).", bold=True, color=INK, height=46)
    para("2. Para actualizar, descarga primero el «Excel actual» desde el panel: ya trae los ID de cada evento. "
         "Edita ese archivo en lugar de empezar de cero.", height=32)
    para("3. En filas nuevas deja ID_EVENTO vacío; el sistema asigna uno (EVT-0001, EVT-0002…). "
         "Nunca cambies ni reutilices un ID existente.", height=32)
    para("4. Fechas en formato dd/mm/aaaa y horas en formato 24 h (08:00, 14:30). La hora es la de Colombia.", height=32)
    row += 1

    title("Columnas")
    table(
        ["Columna", "¿Obligatoria?", "Formato", "Ejemplo", "Notas"],
        [
            ["ID_EVENTO", "No", "Texto", "EVT-0001", "Vacío en filas nuevas. Único: no puede repetirse."],
            ["TIPO", "Sí", "Lista", "SESION_ZAJUNA", "SESION_ZAJUNA = sesión Zajuna, prioritaria (azul intenso). SESION_ADICIONAL = sesión adicional (cian). DUDAS = espacio de dudas por chat (ámbar). Las tres son en vivo con hora de inicio. GRABACION = grabación disponible desde su fecha (rosa). ENTREGA (rojo), CUESTIONARIO (verde) y FORO (violeta) usan hora límite."],
            ["NOMBRE", "Sí", "Texto", "Matemáticas II", "Máximo 150 caracteres."],
            ["FECHA", "Sí", "dd/mm/aaaa", "05/10/2026", "Fecha de la sesión, del espacio de dudas o de la entrega."],
            ["HORA_INICIO", "Sí en sesiones y dudas", "HH:MM (24 h)", "08:00", "En entregas es la hora límite. Si una entrega no tiene hora, vence a las 11:59 PM."],
            ["HORA_FIN", "No", "HH:MM (24 h)", "10:00", "Solo para sesiones y dudas. Debe ser posterior a la hora de inicio. Sin hora fin, se considera de 1 hora."],
            ["LINK", "No", "https://…", "https://meet.google.com/abc", "Sesiones: botón «Ingresar a la sesión». Dudas: botón «Ir al chat de dudas» (agrégalo cuando se tenga). Sin enlace se publica sin botón."],
            ["DESCRIPCION", "No", "Texto", "Traer calculadora", "Máximo 2.000 caracteres."],
            ["INSTRUCTOR", "No", "Texto", "Ana Pérez", "Quién dicta la sesión. En cualquier sesión (SESION_ZAJUNA, SESION_ADICIONAL, DUDAS) y GRABACION (la grabación lleva el instructor de la sesión; puede repetirse el mismo día); en otros tipos se ignora. Máximo 120 caracteres."],
        ],
    )
    row += 1

    title("Ejemplos")
    table(
        ["ID_EVENTO", "TIPO", "NOMBRE", "FECHA · HORAS", "LINK / DESCRIPCIÓN"],
        [
            ["EVT-0001", "SESION_ZAJUNA", "Matemáticas II", "05/10/2026 · 08:00 – 10:00", "https://zajuna.sena.edu.co/… · Unidad 3 · Instructor: Ana Pérez"],
            ["(vacío)", "SESION_ADICIONAL", "Refuerzo de Matemáticas", "06/10/2026 · 18:00 – 19:00", "https://meet.google.com/abc"],
            ["(vacío)", "DUDAS", "Resolución de dudas", "07/10/2026 · 19:00 – 20:00", "Enlace al chat cuando se tenga"],
            ["(vacío)", "GRABACION", "Grabación: Matemáticas II", "05/10/2026 · sin hora", "https://youtu.be/abc · Clase del 5 de octubre · Instructor: Ana Pérez"],
            ["EVT-0002", "ENTREGA", "Taller 2", "06/10/2026 · 23:59", "Subir en PDF"],
            ["(vacío)", "CUESTIONARIO", "Quiz unidad 2", "08/10/2026 · sin hora", "Vence a las 11:59 PM de ese día"],
            ["(vacío)", "FORO", "Foro: la complejidad de la vida", "07/03/2027 · 23:59", "Participación en el foro de la plataforma"],
        ],
        fills=[SES_FILL, ADI_FILL, DUD_FILL, REC_FILL, ENT_FILL, QUIZ_FILL, FORO_FILL],
    )
    row += 1

    title("Errores frecuentes")
    table(
        ["Problema", "", "", "", "Cómo corregirlo"],
        [
            ["Fecha como 35/15/2026", "", "", "", "Revisa día y mes. Debe ser una fecha real en formato dd/mm/aaaa."],
            ["ID_EVENTO repetido", "", "", "", "Cada ID aparece una sola vez. En filas nuevas deja la celda vacía."],
            ["Tipo «REUNION»", "", "", "", "Usa solo SESION_ZAJUNA, SESION_ADICIONAL, DUDAS, GRABACION, ENTREGA, CUESTIONARIO o FORO."],
            ["Sesión o dudas sin hora de inicio", "", "", "", "Las sesiones y los espacios de dudas necesitan HORA_INICIO."],
            ["Hora fin antes del inicio", "", "", "", "Corrige HORA_FIN o déjala vacía."],
            ["Link sin https://", "", "", "", "Copia el enlace completo, empezando por https://"],
        ],
        fills=[SOFT_FILL] * 6,
    )
    for r in range(row - 7, row):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)


def main():
    wb = Workbook()
    build_calendar_sheet(wb.active)
    build_instructions_sheet(wb)
    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Plantilla creada: {OUT}")


if __name__ == "__main__":
    main()
